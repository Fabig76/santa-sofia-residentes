from pathlib import Path
from openpyxl import load_workbook
import re, json

ROOT = Path('/root/santa-sofia-residentes')
errors = []

def ok(cond, msg):
    if not cond:
        errors.append(msg)
    print(('OK  ' if cond else 'ERR ') + msg)

# 1. Archivos requeridos
for rel in ['index.html','assets/styles.css','assets/logo.png','js/app.js','apps-script/Codigo.gs','README.md','docs/GUIA-PROYECTO-SANTA-SOFIA.md','data/matriculas_santa_sofia_COMPLETADO_PROPUESTA.xlsx','data/reporte_completado_matriculas_santa_sofia.csv']:
    ok((ROOT/rel).exists(), f'existe {rel}')

index = (ROOT/'index.html').read_text(encoding='utf-8')
app = (ROOT/'js/app.js').read_text(encoding='utf-8')
gs = (ROOT/'apps-script/Codigo.gs').read_text(encoding='utf-8')
css = (ROOT/'assets/styles.css').read_text(encoding='utf-8')

# 2. Identidad Santa Sofía
for txt in ['Santa Sofía Club Residencial V.I.S','901142051-3','CR 27 # 44-25, Armenia, Quindío','santasofia.clubresidencial@gmail.com']:
    ok(txt in index or txt in gs or txt in app, f'identidad contiene {txt}')

for bad in ['900770444','Bello / Niquía','urb.cerroazul@gmail.com','AKfycbxp','1ceGtZDUJHX4yxs5_ydDwLwtrkOcZwYh09WUG0st_b0Y','16gxeAkcTIWnuwkBFBaHW7Y-nUHaMdtovNzUBaupytPc']:
    ok(bad not in index+app+gs+css, f'sin residuo {bad}')

# 3. Configuración Apps Script
ok("const SHEET_ID = '1xL359rDrhb3_qbhY-tm2MfPXKBqbAehC3zWzsMv1PUo';" in gs, 'SHEET_ID Santa Sofía principal correcto')
ok("const MATRICULAS_SHEET_ID = '1qEnC5BCRags2r_RHQiB0LjK6Or1rx31Gvopr22-n12w';" in gs, 'MATRICULAS_SHEET_ID nativo correcto')
ok("return 'SS-' + String(max + 1).padStart(4, '0');" in gs, 'prefijo SS en getNextFormId')
ok('requiereManual' in gs, 'backend maneja requiereManual')
ok('isMatriculaPendiente' in gs, 'backend detecta X/pendiente')

# 4. Frontend
ok("const APPS_SCRIPT_URL = window.APPS_SCRIPT_URL || '';" in app, 'frontend queda desconectado hasta URL /exec real')
ok('data-required-manual' in app, 'frontend marca campos manuales obligatorios')
ok('Para continuar, debe escribir la matrícula inmobiliaria' in app, 'mensaje bloqueo matrícula manual')
ok('La administración no tiene disponible' in app, 'mensaje aviso administración sin matrícula')

# 5. Excel corregido
xlsx = ROOT/'data/matriculas_santa_sofia_COMPLETADO_PROPUESTA.xlsx'
wb = load_workbook(xlsx, data_only=True, read_only=True)
ok(set(['Resumen y Matrículas','Apartamentos','Parqueaderos']).issubset(set(wb.sheetnames)), 'xlsx contiene hojas esperadas')
apt = wb['Apartamentos']
parq = wb['Parqueaderos']
# Casos esperados
val_311 = None
for i,row in enumerate(apt.iter_rows(values_only=True), start=1):
    if i >= 6 and str(row[2]).strip() == '311':
        val_311 = str(row[7]).strip(); break
ok(val_311 == '280-214831', 'lookup apto 311 esperado 280-214831')
val_111 = None; estado_111 = None
for i,row in enumerate(apt.iter_rows(values_only=True), start=1):
    if i >= 6 and str(row[2]).strip() == '111':
        val_111 = str(row[7]).strip(); estado_111 = str(row[8]).strip(); break
ok('X' in val_111, 'apto 111 sigue pendiente/XXXX')
ok('residente debe digitar manualmente' in estado_111, 'apto 111 marcado manual')
val_m34 = None
val_m60 = None; estado_m60 = None
for i,row in enumerate(parq.iter_rows(values_only=True), start=1):
    if i >= 6 and str(row[2]).strip().upper() == 'MOTO 34':
        val_m34 = str(row[7]).strip()
    if i >= 6 and str(row[2]).strip().upper() == 'MOTO 60':
        val_m60 = str(row[7]).strip(); estado_m60 = str(row[8]).strip()
ok(val_m34 == '280-226102', 'Moto 34 completada 280-226102')
ok('X' in val_m60, 'Moto 60 no completada por ambigua')
ok('AMBIGUA' in estado_m60, 'Moto 60 marcada ambigua')

# 6. Conteo
pending = completed = ambiguous = 0
for ws in [apt, parq]:
    for i,row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 6 or len(row) < 9: continue
        st = str(row[8] or '')
        if st.startswith('PENDIENTE'): pending += 1
        if st.startswith('Completada'): completed += 1
        if st.startswith('AMBIGUA'): ambiguous += 1
ok(completed == 27, f'conteo completadas 27 (actual {completed})')
ok(ambiguous == 1, f'conteo ambiguas 1 (actual {ambiguous})')
ok(pending == 116, f'conteo pendientes 116 (actual {pending})')

if errors:
    print('\nFALLARON', len(errors), 'pruebas')
    raise SystemExit(1)
print('\nTODAS LAS PRUEBAS PASARON')
